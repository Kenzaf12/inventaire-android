# -*- coding: utf-8 -*-
"""
Génère la présentation de soutenance de stage à partir des images réelles
extraites du rapport (logos, diagrammes UML, architectures, captures d'écran).

Pré-requis : pip install python-pptx pillow
Lancer     : python generate_ppt.py
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# ---------------------------------------------------------------- couleurs
DARK   = RGBColor(0x1B, 0x34, 0x4A)
TEAL   = RGBColor(0x26, 0x6F, 0x8E)
GREEN  = RGBColor(0x5A, 0x7D, 0x2B)
LIGHT  = RGBColor(0xF3, 0xF6, 0xF8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREY   = RGBColor(0x55, 0x66, 0x70)
PURPLE = RGBColor(0x6A, 0x3D, 0x9A)

ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ppt_assets")

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def slide():
    return prs.slides.add_slide(BLANK)


def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def rect(s, x, y, w, h, color):
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def txt(s, x, y, w, h, text, size=18, color=DARK, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return tb


def bullets(s, x, y, w, h, items, size=16, color=DARK, gap=6):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = "•  " + it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return tb


def A(stem):
    for f in sorted(os.listdir(ASSETS)):
        if f.startswith(stem + "_"):
            return os.path.join(ASSETS, f)
    return os.path.join(ASSETS, stem + ".png")


def img_fit(s, path, x, y, w, h, align="center", valign="middle"):
    if not os.path.exists(path):
        rect(s, x, y, w, h, LIGHT)
        return
    iw, ih = Image.open(path).size
    if (w / h) < (iw / ih):
        nw, nh = w, int(w * ih / iw)
    else:
        nh, nw = h, int(h * iw / ih)
    nx = x + (w - nw) // 2 if align == "center" else (x if align == "left" else x + w - nw)
    ny = y + (h - nh) // 2 if valign == "middle" else (y if valign == "top" else y + h - nh)
    s.shapes.add_picture(path, nx, ny, nw, nh)


def header(s, title, num=None):
    bg(s, WHITE)
    rect(s, 0, 0, SW, Inches(1.05), TEAL)
    rect(s, 0, Inches(1.05), SW, Inches(0.06), GREEN)
    txt(s, Inches(0.55), 0, Inches(10.8), Inches(1.05), title,
        25, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    img_fit(s, A("p01_3"), Inches(11.7), Inches(0.18), Inches(1.35), Inches(0.7))
    if num is not None:
        txt(s, Inches(12.4), Inches(7.02), Inches(0.8), Inches(0.4),
            str(num), 11, GREY, align=PP_ALIGN.RIGHT)


# ================================================================ 1 — COUVERTURE
s = slide()
bg(s, WHITE)
rect(s, 0, 0, Inches(0.35), SH, TEAL)
rect(s, Inches(0.35), 0, Inches(0.08), SH, GREEN)
img_fit(s, A("p01_0"), Inches(1.3),  Inches(0.30), Inches(1.5), Inches(1.5))
img_fit(s, A("p01_1"), Inches(5.9),  Inches(0.40), Inches(1.5), Inches(1.4))
img_fit(s, A("p01_2"), Inches(9.7),  Inches(0.55), Inches(2.6), Inches(1.1))
txt(s, Inches(1.0), Inches(2.05), Inches(11.5), Inches(0.5),
    "RAPPORT DE STAGE — SOUTENANCE", 15, TEAL, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(1.0), Inches(2.55), Inches(11.5), Inches(1.4),
    "Conception et Développement d'une Application Android\n"
    "Intelligente pour la Gestion de l'Inventaire Physique",
    29, DARK, bold=True, align=PP_ALIGN.CENTER)
img_fit(s, A("p01_3"), Inches(5.4), Inches(4.05), Inches(2.5), Inches(0.95))
txt(s, Inches(1.0), Inches(5.0), Inches(11.5), Inches(0.5),
    "Audit, Ingénierie en Consulting et Formation — Casablanca",
    15, GREY, align=PP_ALIGN.CENTER, italic=True)
rect(s, Inches(1.8), Inches(5.7), Inches(9.7), Inches(1.15), LIGHT)
txt(s, Inches(2.2), Inches(5.82), Inches(5.0), Inches(1.0),
    "Réalisé par :\nKenza FOUDALI", 15, DARK, bold=True)
txt(s, Inches(6.9), Inches(5.82), Inches(4.5), Inches(1.0),
    "Encadrante :\nMme Laila MAADIR", 15, DARK, bold=True)
txt(s, Inches(1.0), Inches(6.95), Inches(11.5), Inches(0.4),
    "Année Universitaire 2025 / 2026   ·   du 04 mai au 13 juin 2026",
    12, GREY, align=PP_ALIGN.CENTER)

# ================================================================ 2 — SOMMAIRE
s = slide()
header(s, "Plan de la présentation", 2)
plan = [
    "1.  Présentation de l'entreprise & cadre du stage",
    "2.  Problématique & objectifs du projet",
    "3.  Déroulement & planification (Gantt / PERT)",
    "4.  Conception & modélisation UML",
    "5.  Architecture du système",
    "6.  Technologies utilisées",
    "7.  Réalisation & démonstration de l'application",
    "8.  Bilan, difficultés & perspectives",
]
rect(s, Inches(1.2), Inches(1.6), Inches(0.12), Inches(5.2), GREEN)
bullets(s, Inches(1.6), Inches(1.75), Inches(10.5), Inches(5.2), plan, size=20, gap=13)

# ================================================================ 3 — INTRODUCTION
s = slide()
header(s, "Introduction & contexte", 3)
txt(s, Inches(0.7), Inches(1.35), Inches(11.9), Inches(1.4),
    "La gestion de l'inventaire physique des immobilisations est une mission "
    "essentielle du métier d'audit de CF Consult. Réalisée manuellement (saisie "
    "Excel, pointage papier), elle reste lente, source d'erreurs et difficile à tracer.",
    17, DARK)
rect(s, Inches(0.7), Inches(3.0), Inches(5.8), Inches(3.7), LIGHT)
txt(s, Inches(1.0), Inches(3.2), Inches(5.3), Inches(0.5), "Problèmes actuels", 18, TEAL, bold=True)
bullets(s, Inches(1.0), Inches(3.8), Inches(5.3), Inches(2.8), [
    "Saisie manuelle longue et fastidieuse",
    "Risque d'erreurs et de doublons",
    "Aucun suivi en temps réel",
    "Pas de responsabilisation par agent",
    "Absence de preuve photo des biens",
], size=15, gap=11)
rect(s, Inches(6.8), Inches(3.0), Inches(5.8), Inches(3.7), TEAL)
txt(s, Inches(7.1), Inches(3.2), Inches(5.3), Inches(0.5), "Solution proposée", 18, WHITE, bold=True)
bullets(s, Inches(7.1), Inches(3.8), Inches(5.3), Inches(2.8), [
    "Application mobile Android native",
    "Scan QR code & code-barres (ZXing)",
    "Reconnaissance d'objet par IA (Gemini)",
    "Photo & validation des biens",
    "Export Excel / PDF & tableau de bord admin",
], size=15, color=WHITE, gap=11)

# ================================================================ 4 — ENTREPRISE
s = slide()
header(s, "Présentation de CF Consult", 4)
img_fit(s, A("p01_3"), Inches(0.8), Inches(1.45), Inches(4.0), Inches(1.5))
rect(s, Inches(0.8), Inches(3.2), Inches(4.2), Inches(0.55), TEAL)
txt(s, Inches(1.0), Inches(3.2), Inches(4.0), Inches(0.55), "Fiche d'identité", 16, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, Inches(0.9), Inches(3.95), Inches(4.3), Inches(2.8), [
    "Activité : Audit, Ingénierie,",
    "    Consulting & Formation",
    "Siège : Casablanca, Maroc",
    "Mission : inventaire physique",
    "    des immobilisations",
], size=14, gap=9)
img_fit(s, A("p17_1"), Inches(5.4), Inches(1.55), Inches(7.2), Inches(2.4))
img_fit(s, A("p18_1"), Inches(6.95), Inches(4.05), Inches(2.9), Inches(2.9))

# ================================================================ 5 — METIER
s = slide()
header(s, "Le métier de l'inventaire physique", 5)
txt(s, Inches(0.7), Inches(1.35), Inches(11.9), Inches(1.0),
    "Recenser, identifier et qualifier l'ensemble des biens d'une organisation, "
    "puis rapprocher le terrain des données comptables.", 17, DARK)
steps = [("1", "Préparer", "Import du fichier\ncomptable (Excel)"),
         ("2", "Recenser", "Scan QR / code-barres\nsur le terrain"),
         ("3", "Qualifier", "État, photo &\nvalidation agent"),
         ("4", "Restituer", "Rapports Excel / PDF\n& écarts")]
x = Inches(0.7)
for n, t, d in steps:
    rect(s, x, Inches(2.9), Inches(2.85), Inches(2.6), LIGHT)
    rect(s, x, Inches(2.9), Inches(2.85), Inches(0.7), TEAL)
    txt(s, x, Inches(2.9), Inches(2.85), Inches(0.7), "Étape " + n, 15, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x, Inches(3.72), Inches(2.85), Inches(0.5), t, 18, DARK, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Emu(x + Inches(0.2)), Inches(4.3), Inches(2.45), Inches(1.1), d, 13, GREY, align=PP_ALIGN.CENTER)
    x = Emu(x + Inches(3.05))
txt(s, Inches(0.7), Inches(5.85), Inches(11.9), Inches(0.9),
    "Objectif : remplacer un processus manuel lent et faillible par un outil "
    "mobile rapide, fiable, traçable et intelligent.", 16, TEAL, bold=True, align=PP_ALIGN.CENTER)

# ================================================================ 6 — OBJECTIFS
s = slide()
header(s, "Objectifs du projet", 6)
rect(s, Inches(0.7), Inches(1.6), Inches(5.85), Inches(4.9), LIGHT)
rect(s, Inches(6.75), Inches(1.6), Inches(5.85), Inches(4.9), LIGHT)
txt(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(0.5), "Objectifs fonctionnels", 18, TEAL, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(6.75), Inches(1.7), Inches(5.85), Inches(0.5), "Objectifs avancés", 18, GREEN, bold=True, align=PP_ALIGN.CENTER)
bullets(s, Inches(1.0), Inches(2.45), Inches(5.3), Inches(4.0), [
    "Authentifier agents & admin (JWT)",
    "Importer les données comptables",
    "Scanner QR codes et codes-barres",
    "Saisir / modifier l'inventaire avec photo",
], size=16, gap=17)
bullets(s, Inches(7.05), Inches(2.45), Inches(5.3), Inches(4.0), [
    "Identifier un bien inconnu par IA",
    "Suivre l'activité des agents",
    "Exporter les rapports Excel & PDF",
    "Tableau de bord administrateur",
], size=16, gap=17)

# ================================================================ 7 — PLANNING
s = slide()
header(s, "Déroulement du stage — 6 semaines", 7)
img_fit(s, A("p21_1"), Inches(0.6), Inches(1.4), Inches(12.1), Inches(5.7))

# ================================================================ 8 — GANTT
s = slide()
header(s, "Planification — Diagramme de Gantt", 8)
img_fit(s, A("p28_0"), Inches(0.4), Inches(1.4), Inches(12.5), Inches(5.7))

# ================================================================ 9 — PERT
s = slide()
header(s, "Planification — Diagramme de PERT", 9)
img_fit(s, A("p29_0"), Inches(0.5), Inches(1.4), Inches(12.3), Inches(5.7))

# ================================================================ 10 — CAS D'UTILISATION
s = slide()
header(s, "Conception — Acteurs & cas d'utilisation", 10)
img_fit(s, A("p33_0"), Inches(3.7), Inches(1.25), Inches(5.9), Inches(6.0))
txt(s, Inches(0.5), Inches(1.5), Inches(3.0), Inches(0.5), "Acteurs", 18, TEAL, bold=True)
bullets(s, Inches(0.5), Inches(2.1), Inches(3.0), Inches(2.8),
        ["Agent : scan, saisie, photo, validation, export",
         "Administrateur : import, gestion des agents, suivi"], size=14, gap=14)
txt(s, Inches(9.8), Inches(1.5), Inches(3.3), Inches(0.5), "Inclusions", 16, GREEN, bold=True)
bullets(s, Inches(9.8), Inches(2.1), Inches(3.3), Inches(2.8),
        ["Scan → modifier un article",
         "Photo → identifier via IA Gemini"], size=14, gap=14)

# ================================================================ 11 — CLASSES
s = slide()
header(s, "Conception — Diagramme de classes", 11)
img_fit(s, A("p35_0"), Inches(3.4), Inches(1.3), Inches(6.6), Inches(5.9))
bullets(s, Inches(0.5), Inches(1.8), Inches(2.7), Inches(5.0),
        ["Agent", "Localisation", "InventEquipement", "InventAutres",
         "Etat", "Station", "Activité"], size=15, gap=17)

# ================================================================ 12 — ARCHITECTURE COMPLETE
s = slide()
header(s, "Architecture complète du système", 12)
img_fit(s, A("p24_1"), Inches(0.6), Inches(1.35), Inches(12.1), Inches(5.9))

# ================================================================ 13 — ARCHITECTURE 3 TIERS
s = slide()
header(s, "Architecture trois tiers", 13)
img_fit(s, A("p30_0"), Inches(2.4), Inches(1.3), Inches(8.5), Inches(5.9))

# ================================================================ 14 — ARCHITECTURE COUCHES
s = slide()
header(s, "Architecture en couches du backend", 14)
img_fit(s, A("p31_0"), Inches(3.6), Inches(1.3), Inches(6.3), Inches(5.9))
bullets(s, Inches(0.5), Inches(1.9), Inches(2.9), Inches(5.0),
        ["Controller (REST)", "Service (métier)", "Repository (JPA)",
         "Modèle (entités)", "Spring Security + JWT"], size=15, gap=19)

# ================================================================ 15 — TECHNOLOGIES
s = slide()
header(s, "Technologies utilisées", 15)
cols = [
    ("Backend", GREEN, ["Java 17", "Spring Boot (API REST)", "Spring Data JPA",
                        "Spring Security + JWT", "Maven"]),
    ("Base de données", TEAL, ["PostgreSQL", "Modèle relationnel", "JDBC", "openpyxl (Excel)"]),
    ("Mobile Android", DARK, ["Java natif", "Retrofit (HTTP)", "ZXing (scan)",
                             "Material Design 3", "FileProvider"]),
    ("IA & Outils", PURPLE, ["Google Gemini 2.5 Flash", "Reconnaissance d'objets",
                            "Git & GitHub", "Postman"]),
]
x = Inches(0.55)
w = Inches(3.0)
for tc, color, items in cols:
    rect(s, x, Inches(1.55), w, Inches(5.3), LIGHT)
    rect(s, x, Inches(1.55), w, Inches(0.75), color)
    txt(s, x, Inches(1.55), w, Inches(0.75), tc, 15, WHITE, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, Emu(x + Inches(0.18)), Inches(2.5), Emu(w - Inches(0.3)), Inches(4.2),
            items, size=14, gap=13)
    x = Emu(x + Inches(3.12))

# ================================================================ 16 — SEQ AUTH
s = slide()
header(s, "Fonctionnement — Authentification JWT", 16)
img_fit(s, A("p36_0"), Inches(3.3), Inches(1.3), Inches(6.7), Inches(5.9))

# ================================================================ 17 — SEQ SCAN
s = slide()
header(s, "Fonctionnement — Scan QR & code-barres", 17)
img_fit(s, A("p37_0"), Inches(2.8), Inches(1.25), Inches(7.7), Inches(6.0))

# ================================================================ 18 — SEQ IA
s = slide()
header(s, "Fonctionnement — Reconnaissance IA (Gemini)", 18)
img_fit(s, A("p39_0"), Inches(3.4), Inches(1.3), Inches(6.5), Inches(5.9))


# ---------------------------------------------------------------- réalisation
def screen_slide(title, num, shots):
    s = slide()
    header(s, title, num)
    n = len(shots)
    total = Inches(2.7) * n + Inches(0.7) * (n - 1)
    x = Emu((SW - total) // 2)
    for path, cap in shots:
        img_fit(s, path, x, Inches(1.4), Inches(2.7), Inches(4.9), valign="top")
        txt(s, x, Inches(6.35), Inches(2.7), Inches(0.6), cap, 14, TEAL, bold=True, align=PP_ALIGN.CENTER)
        x = Emu(x + Inches(2.7) + Inches(0.7))


# ================================================================ 19
screen_slide("Réalisation — Connexion & tableau de bord", 19, [
    (A("p64_1"), "Écran de connexion (JWT)"),
    (A("p65_1"), "Tableau de bord principal"),
])
# ================================================================ 20
screen_slide("Réalisation — Liste & fiche d'un bien", 20, [
    (A("p66_1"), "Liste & recherche des biens"),
    (A("p67_1"), "Fiche d'inventaire d'un bien"),
])
# ================================================================ 21 — IA
s = slide()
header(s, "Réalisation — Reconnaissance par IA", 21)
img_fit(s, A("p68_1"), Inches(1.2), Inches(1.4), Inches(3.0), Inches(5.6), valign="top")
txt(s, Inches(4.7), Inches(1.9), Inches(8.0), Inches(0.6),
    "Identification d'un bien inconnu par photo", 20, TEAL, bold=True)
bullets(s, Inches(4.7), Inches(2.8), Inches(8.0), Inches(3.6), [
    "L'agent prend une photo du bien (caméra ou galerie)",
    "L'image est envoyée au backend puis à l'API Gemini 2.5 Flash",
    "L'IA retourne désignation, type, marque et état estimé",
    "Les champs de la fiche sont pré-remplis automatiquement",
    "Gestion des erreurs : 3 tentatives en cas de surcharge",
], size=16, gap=15)
# ================================================================ 22
screen_slide("Réalisation — Export & administration", 22, [
    (A("p69_1"), "États & export Excel / PDF"),
    (A("p70_1"), "Tableau de bord administrateur"),
])

# ================================================================ 23 — BILAN
s = slide()
header(s, "Bilan du stage", 23)
rect(s, Inches(0.7), Inches(1.6), Inches(5.85), Inches(4.9), LIGHT)
txt(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(0.5), "Compétences techniques", 18, TEAL, bold=True, align=PP_ALIGN.CENTER)
bullets(s, Inches(1.0), Inches(2.45), Inches(5.3), Inches(4.0), [
    "Développement Android natif (Java)",
    "API REST avec Spring Boot",
    "Sécurité JWT & Spring Security",
    "Base PostgreSQL / JPA",
    "Intégration d'une IA (Gemini)",
    "Architecture trois tiers en couches",
], size=15, gap=13)
rect(s, Inches(6.75), Inches(1.6), Inches(5.85), Inches(4.9), TEAL)
txt(s, Inches(6.75), Inches(1.7), Inches(5.85), Inches(0.5), "Compétences humaines", 18, WHITE, bold=True, align=PP_ALIGN.CENTER)
bullets(s, Inches(7.05), Inches(2.45), Inches(5.3), Inches(4.0), [
    "Analyse d'un besoin métier réel",
    "Gestion et planification de projet",
    "Autonomie et organisation",
    "Communication avec l'encadrante",
    "Rédaction technique & documentation",
    "Résolution de problèmes concrets",
], size=15, color=WHITE, gap=13)

# ================================================================ 24 — PERSPECTIVES
s = slide()
header(s, "Difficultés & perspectives", 24)
txt(s, Inches(0.7), Inches(1.45), Inches(11.9), Inches(0.5), "Difficultés rencontrées", 18, GREEN, bold=True)
bullets(s, Inches(1.0), Inches(2.05), Inches(11.5), Inches(2.0), [
    "Gestion des permissions caméra & du partage de fichiers (FileProvider)",
    "Robustesse des appels à l'API Gemini (gestion des erreurs et reprises)",
    "Compatibilité du stockage selon les versions d'Android",
], size=16, gap=10)
txt(s, Inches(0.7), Inches(4.1), Inches(11.9), Inches(0.5), "Perspectives d'évolution", 18, TEAL, bold=True)
bullets(s, Inches(1.0), Inches(4.7), Inches(11.5), Inches(2.2), [
    "Mode hors-ligne avec synchronisation différée",
    "Génération et impression d'étiquettes QR depuis l'app",
    "Statistiques et tableaux de bord analytiques avancés",
    "Application multi-entreprises et multi-sites",
], size=16, gap=10)

# ================================================================ 25 — CONCLUSION
s = slide()
header(s, "Conclusion", 25)
txt(s, Inches(0.9), Inches(1.8), Inches(11.5), Inches(2.5),
    "Ce stage a permis de concevoir et développer une solution complète et "
    "fonctionnelle de gestion d'inventaire physique pour CF Consult, combinant "
    "une application mobile Android, un backend Spring Boot sécurisé et une "
    "intelligence artificielle de reconnaissance d'objets.", 19, DARK)
txt(s, Inches(0.9), Inches(4.2), Inches(11.5), Inches(2.0),
    "La solution transforme un processus manuel, lent et faillible en un outil "
    "rapide, fiable et traçable — une expérience à la fois technique et "
    "professionnelle très enrichissante.", 19, TEAL, bold=True)

# ================================================================ 26 — MERCI
s = slide()
bg(s, DARK)
rect(s, 0, Inches(3.0), SW, Inches(1.5), TEAL)
txt(s, 0, Inches(2.9), SW, Inches(1.0), "Merci de votre attention", 40, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, 0, Inches(3.95), SW, Inches(0.7), "Avez-vous des questions ?", 22, WHITE, align=PP_ALIGN.CENTER)
txt(s, 0, Inches(6.6), SW, Inches(0.5), "Kenza FOUDALI  ·  CF Consult  ·  2025 / 2026", 14, RGBColor(0xBB,0xCC,0xD5), align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- save
out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "soutenance_stage_inventaire_cfc.pptx")
prs.save(out)
print(f"OK  {out}  ({len(prs.slides._sldIdLst)} slides)")
