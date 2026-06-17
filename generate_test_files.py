"""
Génère les fichiers de test pour l'app Inventaire CFC :
  - equipements_test.xlsx   → à importer dans l'app via Admin > Import Équipements
  - autres_test.xlsx        → à importer dans l'app via Admin > Import Autres
  - qrcodes/                → images PNG des QR codes à scanner

INSTALLATION :
    pip install openpyxl qrcode[pil]

UTILISATION :
    python generate_test_files.py
"""

import os
import json

# ── Excel ────────────────────────────────────────────────────────────────────

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    EQUIP_HEADER_COLOR = "1F4E79"
    AUTRES_HEADER_COLOR = "375623"

    def style_header(ws, color):
        fill = PatternFill("solid", fgColor=color)
        font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center")

    # ── Équipements ──────────────────────────────────────────────────────────
    wb_equip = openpyxl.Workbook()
    ws = wb_equip.active
    ws.title = "Équipements"

    headers = ["refImmo", "cab", "station", "article", "equipement",
               "designation", "modele", "marque", "nserie", "qte",
               "etat", "descTech", "observation"]
    ws.append(headers)

    rows = [
        ["IMM-010", "CAB010", "Casablanca", "ART-010", "Ordinateur Portable",
         "Lenovo ThinkPad X1 Carbon", "X1 Carbon", "Lenovo", "SN-LP-010", 1,
         "En activité", "Intel Core i5, 8GB RAM, 256GB SSD", ""],
        ["IMM-011", "CAB011", "Casablanca", "ART-011", "Ordinateur Portable",
         "Lenovo ThinkPad X1 Carbon", "X1 Carbon", "Lenovo", "SN-LP-011", 1,
         "En activité", "Intel Core i5, 8GB RAM, 256GB SSD", "Batterie à vérifier"],
        ["IMM-012", "CAB012", "Rabat", "ART-012", "Écran",
         "Dell P2422H 24 pouces", "P2422H", "Dell", "SN-EC-012", 2,
         "En activité", "24 pouces Full HD IPS", ""],
        ["IMM-013", "CAB013", "Rabat", "ART-013", "Imprimante",
         "Canon PIXMA G3420", "G3420", "Canon", "SN-IM-013", 1,
         "En activité", "Jet d'encre couleur, WiFi", ""],
        ["IMM-014", "CAB014", "Casablanca", "ART-014", "Téléphone IP",
         "Cisco IP Phone 7945G", "7945G", "Cisco", "SN-TL-014", 3,
         "En activité", "Double ligne, écran couleur", ""],
        ["IMM-015", "CAB015", "Marrakech", "ART-015", "Onduleur",
         "APC Smart-UPS 1500VA", "1500VA", "APC", "SN-ON-015", 1,
         "Hors service", "1500VA / 980W", "Batteries HS - à remplacer"],
        ["IMM-016", "CAB016", "Marrakech", "ART-016", "Scanner",
         "Fujitsu ScanSnap iX1600", "iX1600", "Fujitsu", "SN-SC-016", 1,
         "En activité", "Scanner documentaire WiFi, 40ppm", ""],
    ]
    for r in rows:
        ws.append(r)

    style_header(ws, EQUIP_HEADER_COLOR)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20

    wb_equip.save("equipements_test.xlsx")
    print("✅  equipements_test.xlsx créé (7 lignes)")

    # ── Autres ───────────────────────────────────────────────────────────────
    wb_autres = openpyxl.Workbook()
    ws2 = wb_autres.active
    ws2.title = "Autres"

    headers2 = ["numRef", "refImmo", "cab", "nlocal", "designation",
                "unite", "qte", "montGl", "modele", "marque",
                "fournisseur", "descTech", "etat", "observation"]
    ws2.append(headers2)

    rows2 = [
        ["REF-A010", "IMM-A010", "CAB-A010", "Bureau 201", "Chaise ergonomique",
         "Unité", 6, 4800, "Leap V2", "Steelcase", "Bureau+", "Chaise de bureau ergonomique", "En activité", ""],
        ["REF-A011", "IMM-A011", "CAB-A011", "Bureau 201", "Bureau réglable",
         "Unité", 3, 9000, "Flexispot E7", "Flexispot", "Ergonomia", "Bureau assis-debout motorisé", "En activité", ""],
        ["REF-A012", "IMM-A012", "CAB-A012", "Salle réunion 01", "Tableau blanc",
         "Unité", 2, 1200, "Classic 100x200", "Nobo", "Fournitures Plus", "Tableau blanc magnétique", "En activité", ""],
        ["REF-A013", "IMM-A013", "CAB-A013", "Accueil", "Canapé 3 places",
         "Unité", 1, 5500, "KIVIK", "IKEA", "IKEA Maroc", "Canapé tissu gris", "En activité", "Légèrement usé"],
    ]
    for r in rows2:
        ws2.append(r)

    style_header(ws2, AUTRES_HEADER_COLOR)
    for col in ws2.columns:
        ws2.column_dimensions[col[0].column_letter].width = 18

    wb_autres.save("autres_test.xlsx")
    print("✅  autres_test.xlsx créé (4 lignes)")

except ImportError:
    print("⚠️  openpyxl non installé — lance : pip install openpyxl")

# ── QR Codes ─────────────────────────────────────────────────────────────────

try:
    import qrcode

    os.makedirs("qrcodes", exist_ok=True)

    # QR codes JSON (remplis les champs automatiquement)
    qr_items = [
        # Équipements existants dans le mock server
        {"filename": "qr_CAB001_Ordinateur.png",
         "data": json.dumps({"type": "EQUIPEMENT", "id": 1, "cab": "CAB001",
                             "ref_immo": "IMM-001", "designation": "Dell Latitude 5520",
                             "equipement": "Ordinateur Portable", "marque": "Dell",
                             "modele": "5520", "nserie": "SN123456",
                             "station": "Casablanca", "article": "ART-001", "etat": "En activité"})},
        {"filename": "qr_CAB002_Imprimante.png",
         "data": json.dumps({"type": "EQUIPEMENT", "id": 2, "cab": "CAB002",
                             "ref_immo": "IMM-002", "designation": "HP LaserJet Pro M404dn",
                             "equipement": "Imprimante", "marque": "HP",
                             "modele": "M404dn", "nserie": "SN654321",
                             "station": "Casablanca", "article": "ART-002", "etat": "En activité"})},
        {"filename": "qr_CAB003_Ecran.png",
         "data": json.dumps({"type": "EQUIPEMENT", "id": 3, "cab": "CAB003",
                             "ref_immo": "IMM-003", "designation": "LG 27UL500 4K",
                             "equipement": "Écran", "marque": "LG",
                             "modele": "27UL500", "nserie": "SN789012",
                             "station": "Rabat", "article": "ART-003", "etat": "En activité"})},
        # Autres existants
        {"filename": "qr_CABA001_Chaise.png",
         "data": json.dumps({"type": "AUTRES", "id": 1, "cab": "CAB-A001",
                             "designation": "Chaise de bureau", "marque": "Steelcase",
                             "nlocal": "Bureau 101", "etat": "En activité"})},
        # Code-barres simple (juste le CAB) → cherche dans l'API
        {"filename": "barcode_CAB004_Serveur.png",   "data": "CAB004"},
        {"filename": "barcode_CAB005_Switch.png",    "data": "CAB005"},
        {"filename": "barcode_CABA002_Table.png",    "data": "CAB-A002"},
    ]

    for item in qr_items:
        qr = qrcode.QRCode(box_size=8, border=2)
        qr.add_data(item["data"])
        qr.make(fit=True)
        img = qr.make_image(fill_color="black", back_color="white")
        path = os.path.join("qrcodes", item["filename"])
        img.save(path)
        print(f"✅  qrcodes/{item['filename']}")

    print()
    print("📋  Comment tester le scan :")
    print("    1. Ouvre un fichier PNG depuis qrcodes/ sur ton PC")
    print("    2. Dans l'app, appuie sur le bouton 📷 en bas à droite (Scanner)")
    print("    3. Pointe la caméra du téléphone vers l'écran de ton PC")
    print("    4. L'app ouvre automatiquement la fiche de l'équipement/autre")

except ImportError:
    print("⚠️  qrcode non installé — lance : pip install qrcode[pil]")

print()
print("=" * 55)
print("  Fichiers générés dans :", os.getcwd())
print("=" * 55)
